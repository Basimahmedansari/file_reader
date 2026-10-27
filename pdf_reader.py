
import io
from pathlib import Path

import fitz
import pandas as pd
import pytesseract

from PIL import Image
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def read_pdf(file):
    data = file.getvalue() if hasattr(file, "getvalue") else file.read()
    reader = PdfReader(io.BytesIO(data))

    # Extract text and OCR individual pages when needed.
    pages_text = []

    with fitz.open(stream=data, filetype="pdf") as pdf:
        for index, page in enumerate(pdf):
            text = reader.pages[index].extract_text() or ""

            if not text.strip():
                pix = page.get_pixmap(
                    matrix=fitz.Matrix(2, 2),
                    alpha=False
                )
                image = Image.open(
                    io.BytesIO(pix.tobytes("png"))
                )
                text = pytesseract.image_to_string(image)

            pages_text.append(text)

    return "\n\n".join(pages_text)


def extract_file(file):
    ext = Path(file.name).suffix.lower()
    data = file.getvalue()

    if ext == ".pdf":
        return read_pdf(io.BytesIO(data))

    if ext in {
        ".txt", ".md", ".py", ".json", ".xml", ".html",
        ".css", ".js", ".yaml", ".yml", ".log", ".ini", ".sql"
    }:
        return data.decode("utf-8", errors="replace")

    if ext == ".csv":
        return pd.read_csv(io.BytesIO(data)).to_string(index=False)

    if ext in {".xlsx", ".xls"}:
        sheets = pd.read_excel(
            io.BytesIO(data), sheet_name=None
        )
        return "\n\n".join(
            f"Sheet: {name}\n{df.to_string(index=False)}"
            for name, df in sheets.items()
        )

    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts)

    if ext == ".pptx":
        presentation = Presentation(io.BytesIO(data))
        parts = []

        for number, slide in enumerate(presentation.slides, 1):
            parts.append(f"Slide {number}")

            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text)

                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(cell.text for cell in row.cells)
                        )

        return "\n".join(parts)

    if ext in {
        ".jpg", ".jpeg", ".png", ".webp",
        ".bmp", ".tiff", ".tif"
    }:
        image = Image.open(io.BytesIO(data))
        return pytesseract.image_to_string(image)

    raise ValueError(f"Unsupported file type: {ext}")
